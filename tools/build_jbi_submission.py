"""Build a Journal of Biomedical Informatics submission package."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / "manuscript" / "manuscript.md"
ZH_SOURCE = PROJECT / "manuscript" / "manuscript_zh.md"
SUPPLEMENT = PROJECT / "manuscript" / "supplement.md"
FIGURES = PROJECT / "figures"
OUTPUT = PROJECT / "jbi_submission"


def _extract(text: str, start: str, end: str | None = None) -> str:
    start_index = text.index(start)
    tail = text[start_index:]
    if end and end in tail:
        tail = tail[: tail.index(end)]
    return tail.strip()


def _strip_tables(markdown: str) -> str:
    lines = markdown.splitlines()
    output: list[str] = []
    in_table = False
    for line in lines:
        if line.strip().startswith("|"):
            in_table = True
            continue
        if in_table and not line.strip():
            in_table = False
            output.append("")
            continue
        if not in_table:
            output.append(line)
    return "\n".join(output).strip()


def _renumber_headings(markdown: str, base: str, start: int) -> str:
    output: list[str] = []
    subsection = 0
    for line in markdown.splitlines():
        if line.startswith("### "):
            subsection += 1
            output.append(f"### {start}.{subsection} {line[4:].strip()}")
        elif line.startswith(f"## {base}"):
            output.append(f"## {start} {base}")
        else:
            output.append(line)
    return "\n".join(output)


def _word_count(markdown: str) -> int:
    return len(re.findall(r"\b[\w'’\-]+\b", markdown))


def _remove_title_border(style) -> None:
    ppr = style.element.get_or_add_pPr()
    for existing in list(ppr.findall(qn("w:pBdr"))):
        ppr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    ppr.append(borders)


def _set_table_borders(table, color: str = "D9D9D9") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.append(borders)


def _set_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, width in (
        ("top", "80"),
        ("left", "100"),
        ("bottom", "80"),
        ("right", "100"),
    ):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), width)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)


def _set_table_row_properties(table) -> None:
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.find(qn("w:trPr"))
        if tr_pr is None:
            tr_pr = OxmlElement("w:trPr")
            row._tr.insert(0, tr_pr)
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            tr_pr.append(repeat_header)


def _add_table(
    document: Document,
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 9,
    center_columns: set[int] | None = None,
) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    table.style = "Table Grid"
    _set_table_borders(table)
    _set_cell_margins(table)
    _set_table_row_properties(table)
    if center_columns is None:
        center_columns = set(range(1, len(rows[0])))
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_index in center_columns
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(_normalize_inline_math(str(value)))
            run.font.size = Pt(font_size)
            if row_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "334E68")
                cell._tc.get_or_add_tcPr().append(shading)
            elif row_index % 2 == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F4F7")
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph()


def _table_widths(rows: list[list[str]], available: float = 6.7) -> list[float]:
    column_count = len(rows[0])
    weights = []
    for column_index in range(column_count):
        max_word_length = max(
            len(word)
            for row in rows
            for word in re.split(r"\s+", str(row[column_index]).strip())
            if word
        )
        weights.append(float(min(max(max_word_length + 2, 8), 30)))
    total = sum(weights)
    widths = [available * weight / total for weight in weights]
    return widths


def _table_font_size(rows: list[list[str]]) -> float:
    column_count = len(rows[0])
    if column_count >= 9:
        return 6.8
    if column_count >= 8:
        return 7.2
    if column_count >= 7:
        return 7.7
    if column_count >= 6:
        return 8.2
    longest = max(len(cell) for row in rows for cell in row)
    return 8.5 if longest > 60 else 9


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _normalize_inline_math(text: str) -> str:
    def replace(match: re.Match[str]) -> str:
        expression = match.group(1)
        expression = re.sub(
            r"\\times\s*10\^\{?([+-]?\d+)\}?",
            r" x 10^\1",
            expression,
        )
        return expression.replace("\\", "")

    return re.sub(r"\$([^$]+)\$", replace, text)


def _add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", _normalize_inline_math(text))
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    _remove_title_border(styles["Title"])


def _add_markdown(document: Document, markdown: str) -> None:
    paragraph_lines: list[str] = []

    def flush() -> None:
        if not paragraph_lines:
            return
        text = " ".join(paragraph_lines).strip()
        paragraph_lines.clear()
        if not text:
            return
        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, text)

    lines = markdown.splitlines()
    line_index = 0
    while line_index < len(lines):
        line = lines[line_index]
        stripped = line.strip()
        if not stripped:
            flush()
            line_index += 1
            continue
        if (
            stripped.startswith("|")
            and line_index + 1 < len(lines)
            and re.fullmatch(r"\|?[\s:|-]+\|?", lines[line_index + 1].strip())
        ):
            flush()
            table_lines = []
            while line_index < len(lines) and lines[line_index].strip().startswith("|"):
                table_lines.append(lines[line_index])
                line_index += 1
            rows = [_split_table_row(row) for row in table_lines]
            rows = [rows[0], *rows[2:]]
            longest = max(len(cell) for row in rows for cell in row)
            _add_table(
                document,
                rows,
                _table_widths(rows),
                font_size=_table_font_size(rows),
                center_columns=set() if longest > 60 else None,
            )
            continue
        if stripped.startswith("# "):
            flush()
            paragraph = document.add_heading(stripped[2:].strip(), level=0)
            paragraph.paragraph_format.keep_with_next = True
        elif stripped.startswith("## "):
            flush()
            paragraph = document.add_heading(stripped[3:].strip(), level=1)
            paragraph.paragraph_format.keep_with_next = True
        elif stripped.startswith("### "):
            flush()
            paragraph = document.add_heading(stripped[4:].strip(), level=2)
            paragraph.paragraph_format.keep_with_next = True
        elif stripped.startswith("- "):
            flush()
            paragraph = document.add_paragraph(stripped[2:], style="List Bullet")
            _add_inline_runs(paragraph, paragraph.text)
            paragraph.text = ""
            _add_inline_runs(paragraph, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            flush()
            item_lines = [stripped]
            lookahead = line_index + 1
            while lookahead < len(lines):
                continuation = lines[lookahead]
                continuation_text = continuation.strip()
                if not continuation_text or not continuation.startswith((" ", "\t")):
                    break
                item_lines.append(continuation_text)
                lookahead += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.34)
            paragraph.paragraph_format.first_line_indent = Inches(-0.34)
            paragraph.paragraph_format.space_after = Pt(3)
            _add_inline_runs(paragraph, " ".join(item_lines))
            line_index = lookahead
            continue
        else:
            paragraph_lines.append(stripped)
        line_index += 1
    flush()


def _write_docx(path: Path, markdown: str, title: str) -> None:
    document = Document()
    _style_document(document)
    document.add_heading(title, level=0)
    lines = markdown.lstrip().splitlines()
    if lines and lines[0].startswith("# "):
        heading = lines[0][2:].strip()
        if heading.casefold() == title.casefold():
            lines = lines[1:]
    markdown = "\n".join(lines).lstrip()
    _add_markdown(document, markdown)
    document.save(path)


def _graphical_abstract() -> None:
    width, height = 1328, 531
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    regular = ImageFont.truetype(str(regular_path), 18)
    small = ImageFont.truetype(str(regular_path), 15)
    bold = ImageFont.truetype(str(bold_path), 28)
    box_font = ImageFont.truetype(str(bold_path), 16)
    footer_font = ImageFont.truetype(str(bold_path), 15)

    def center_text(text: str, y: int, font: ImageFont.FreeTypeFont, fill: str) -> None:
        box = draw.textbbox((0, 0), text, font=font)
        text_width = box[2] - box[0]
        draw.text(((width - text_width) / 2, y), text, font=font, fill=fill)

    center_text("Module-level agreement is not replication", 22, bold, "#1F2937")
    center_text(
        "A cross-cohort reproducibility audit of public human IVD transcriptomes",
        62,
        regular,
        "#334E68",
    )
    boxes = [
        (48, 140, 320, 280, "6 public IVD cohorts", "#E6F4F1", "#1F6F78"),
        (
            356,
            140,
            628,
            280,
            "Matched AF I-IV\nGSE70362 vs GSE23130",
            "#E8F0FA",
            "#2F5D8A",
        ),
        (
            664,
            140,
            936,
            280,
            "ECM module significant\nin one cohort",
            "#FFF1E6",
            "#B45309",
        ),
        (
            972,
            140,
            1280,
            280,
            "Gene-level concordance\nbelow 50%",
            "#FDE8E8",
            "#B42318",
        ),
    ]
    for left, top, right, bottom, text, fill, edge in boxes:
        draw.rounded_rectangle(
            (left, top, right, bottom),
            radius=14,
            fill=fill,
            outline=edge,
            width=2,
        )
        lines = text.splitlines()
        line_heights = [
            draw.textbbox((0, 0), line, font=box_font)[3]
            - draw.textbbox((0, 0), line, font=box_font)[1]
            for line in lines
        ]
        total_height = sum(line_heights) + 6 * (len(lines) - 1)
        y = top + (bottom - top - total_height) / 2
        for line in lines:
            box = draw.textbbox((0, 0), line, font=box_font)
            text_width = box[2] - box[0]
            draw.text(
                (left + (right - left - text_width) / 2, y),
                line,
                font=box_font,
                fill=edge,
            )
            y += (box[3] - box[1]) + 6
    arrow_y = 210
    draw.line((328, arrow_y, 350, arrow_y), fill="#6B7280", width=3)
    draw.polygon(
        [(350, arrow_y), (340, arrow_y - 6), (340, arrow_y + 6)],
        fill="#6B7280",
    )
    draw.line((636, arrow_y, 658, arrow_y), fill="#6B7280", width=3)
    draw.polygon(
        [(658, arrow_y), (648, arrow_y - 6), (648, arrow_y + 6)],
        fill="#6B7280",
    )
    draw.line((944, arrow_y, 966, arrow_y), fill="#6B7280", width=3)
    draw.polygon(
        [(966, arrow_y), (956, arrow_y - 6), (956, arrow_y + 6)],
        fill="#6B7280",
    )
    center_text(
        "Claim ladder: exploratory -> internally stable -> contextual concordance -> candidate replication -> FDR-controlled replication",
        328,
        small,
        "#1F2937",
    )
    center_text(
        "No module reached strong cross-cohort replication; report gene-level sign concordance with module direction.",
        420,
        footer_font,
        "#B42318",
    )
    image.save(OUTPUT / "JBI_Graphical_Abstract.png", dpi=(300, 300))
    image.save(OUTPUT / "JBI_Graphical_Abstract.tiff", dpi=(300, 300))


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    source = SOURCE.read_text(encoding="utf-8")
    introduction = _extract(source, "## Introduction", "## Methods")
    methods = _extract(source, "## Methods", "## Results")
    results = _strip_tables(_extract(source, "## Results", "## Discussion"))
    # Strip discussion tables EXCEPT the Box 1 claim-ladder table, which is a
    # main-text display item (counts as the third main table).
    discussion_raw = _extract(source, "## Discussion", "## Data availability")
    box1_match = re.search(r"\*\*Box 1\.", discussion_raw)
    if box1_match:
        discussion = _strip_tables(discussion_raw[: box1_match.start()]) + "\n\n" + discussion_raw[box1_match.start() :].strip()
    else:
        discussion = _strip_tables(discussion_raw)
    data_availability = _extract(source, "## Data availability", "## Figure legends")
    figure_legends = _extract(source, "## Figure legends", "## References")
    references = _extract(source, "## References")

    abstract = """## Structured Abstract

**Objective.** To determine whether module-level significance or direction agreement establishes gene-level cross-cohort replication in public human intervertebral disc transcriptomes.

**Methods.** We audited six public human IVD transcriptome cohorts and used GSE70362 and GSE23130 for a matched annulus fibrosus grade I-IV comparison. Effects were estimated with gene-wise linear models, then assessed using exact sign tests, stratified permutation, empirical null models, split-half reliability, donor-level resampling and expression-matched sensitivity analyses. We introduced a five-level claim ladder ranging from exploratory module signals to FDR-controlled gene-level replication and tested it on seven prespecified mechanisms plus 1,554 Hallmark, KEGG and Reactome modules.

**Results.** The ECM remodelling module was significant in GSE23130 (77.4% positive, p=0.0017) but not in GSE70362 (46.7%, p=0.7077). Gene-level concordance was 23.3% (two-sided p=0.0052); probe and scale sensitivity preserved concordance below 50% but did not preserve below-chance significance. The ECM module did not exceed chance in cross-cohort gene-level agreement. Across 1,554 modules, 673 were significant in one cohort and 31 in both; 14 reached the nominal gene-level gate but no module survived strict FDR. An alias-aware sensitivity and collection-local FDR analysis preserved the qualitative result.

**Conclusion.** Module-level significance does not establish replication. Public-data IVD claims should report module direction, gene-level sign concordance, internal reliability, matched comparability and multiplicity control before using replication language.

**Keywords:** reproducibility; transcriptomics; bioinformatics; gene-sets; intervertebral-disc
"""

    related_work = """## 2 Related Work

Reproducibility studies have shown that original effect sizes often shrink on replication [3,4,5], and random or weakly related gene signatures can be significantly associated with clinical outcomes [19]. Marginal gene selection can also be biased by coexpression patterns [6]. At the assay level, batch and latent variation affect high-throughput measurements [15,16], and cross-platform concordance depends on analysis and probe choices [18]. Reporting frameworks address adjacent problems of transparent study description and reusable data and code [20,21]. The contribution of this audit is not a new statistic or algorithm. It integrates a matched IVD comparison, a gene-level reproducibility gate and sensitivity analyses for donor dependence, probe selection, gene identity, scale and FDR-family structure.
"""

    statement = """## Statement of Significance

| Problem or Issue | What is Already Known | What this Paper Adds | Who would Benefit |
|---|---|---|---|
| Module-level significance and direction agreement are often presented as evidence of replication in public-data transcriptomics. | Random signatures may be significant, batch and platform effects can distort cross-cohort comparisons, and gene-set aggregation can hide gene-level disagreement. | A matched IVD audit separates module significance, internal reliability, contextual consistency, gene-level concordance and FDR-controlled replication, with sensitivity analyses across the full analysis chain. | Authors, reviewers, editors and methodologists evaluating public transcriptome claims. |
"""

    conclusion = """## 6 Conclusion

Public human IVD transcriptomes can support strong module-level signals even when the same genes do not replicate across matched cohorts. The claim ladder and sensitivity workflow provide a practical reporting standard for distinguishing exploratory module findings from gene-level replication. The approach is directly applicable to other public human transcriptome audits with ordered severity or clinical endpoints.
"""

    table1_md = """**Table 1. Cohort diagnostics.**

| Analysis set | Samples | Donors | Genes | Up | Median effect | Shift |
|---|---|---|---|---|---|---|
| GSE70362 AF, all grades | 24 | 18 | 18,479 | 47.6% | -0.0027 | No |
| GSE70362 AF, I-IV | 20 | 16 | 18,479 | 43.6% | -0.0089 | No |
| GSE23130, all processing | 23 | NA | 22,107 | 49.7% | -0.0002 | No |
| GSE23130 LCM, I-IV | 15 | NA | 22,107 | 55.0% | +0.0075 | No |
"""

    table2_md = """**Table 2. Primary ECM comparison and probe-collapsing sensitivity.**

| Comparison | GSE70362 up | GSE23130 up | Concordance | Two-sided p |
|---|---|---|---|---|
| All available AF/LCM | 46.7% | 87.1% | 40.0% | 0.3616 |
| Strict I-IV, adjusted (primary) | 46.7% | 77.4% | 23.3% | 0.0052 |
| Strict I-IV, grade only | 50.0% | 74.2% | 23.3% | 0.0052 |
| First-probe collapsing | 46.7% | 83.9% | 30.0% | 0.0428 |
| Median-probe collapsing | 46.7% | 74.2% | 33.3% | 0.0987 |
| Random-probe, median (2.5-97.5%) | 46.7% | 71.0% (61.3-80.6%) | 36.7% (23.3-48.4%) | 0.2005 |
"""

    manuscript_md = "\n\n".join(
        [
            "# Module-level agreement is not replication: a reporting and reproducibility gate for public human intervertebral disc transcriptomes",
            "Article type: Special Communication",
            abstract,
            "## 1 Introduction",
            introduction.replace("## Introduction", "").strip(),
            statement,
            related_work,
            _renumber_headings(methods, "Methods", 3),
            _renumber_headings(results, "Results", 4) + "\n\n" + table1_md + "\n\n" + table2_md,
            _renumber_headings(discussion, "Discussion", 5),
            conclusion,
            data_availability.replace("## Data availability", "## Data Availability", 1),
            figure_legends.replace("## Figure legends", "## Figure Legends", 1),
            references,
        ]
    )
    (OUTPUT / "JBI_Manuscript.md").write_text(manuscript_md, encoding="utf-8")
    body_markdown = manuscript_md.split("## 1 Introduction", 1)[1].split(
        "## Data Availability", 1
    )[0]
    body_word_count = _word_count(_strip_tables(body_markdown))

    document = Document()
    _style_document(document)
    document.add_heading(
        "Module-level agreement is not replication: a reporting and reproducibility gate for public human intervertebral disc transcriptomes",
        level=0,
    )
    document.add_paragraph("Article type: Special Communication")
    _add_markdown(document, abstract)
    _add_markdown(
        document,
        "## 1 Introduction\n\n" + introduction.replace("## Introduction", "").strip(),
    )
    _add_markdown(document, statement)
    _add_markdown(document, related_work)
    _add_markdown(document, _renumber_headings(methods, "Methods", 3))
    _add_markdown(document, _renumber_headings(results, "Results", 4))
    _add_markdown(document, "**Table 1. Cohort diagnostics.**")
    _add_table(
        document,
        [
            ["Analysis set", "Samples", "Donors", "Genes", "Up", "Median effect", "Shift"],
            ["GSE70362 AF, all grades", "24", "18", "18,479", "47.6%", "-0.0027", "No"],
            ["GSE70362 AF, I-IV", "20", "16", "18,479", "43.6%", "-0.0089", "No"],
            ["GSE23130, all processing", "23", "NA", "22,107", "49.7%", "-0.0002", "No"],
            ["GSE23130 LCM, I-IV", "15", "NA", "22,107", "55.0%", "+0.0075", "No"],
        ],
        [2.2, 0.72, 0.72, 0.83, 0.68, 1.0, 0.65],
    )
    _add_markdown(document, "**Table 2. Primary ECM comparison and probe-collapsing sensitivity.**")
    _add_table(
        document,
        [
            ["Comparison", "GSE70362 up", "GSE23130 up", "Concordance", "Two-sided p"],
            ["All available AF/LCM", "46.7%", "87.1%", "40.0%", "0.3616"],
            ["Strict I-IV, adjusted (primary)", "46.7%", "77.4%", "23.3%", "0.0052"],
            ["Strict I-IV, grade only", "50.0%", "74.2%", "23.3%", "0.0052"],
            ["First-probe collapsing", "46.7%", "83.9%", "30.0%", "0.0428"],
            ["Median-probe collapsing", "46.7%", "74.2%", "33.3%", "0.0987"],
            ["Random-probe, median (2.5-97.5%)", "46.7%", "71.0% (61.3-80.6%)", "36.7% (23.3-48.4%)", "0.2005"],
        ],
        [2.35, 1.0, 1.0, 0.95, 0.9],
    )
    _add_markdown(document, _renumber_headings(discussion, "Discussion", 5))
    _add_markdown(document, conclusion)
    _add_markdown(
        document,
        data_availability.replace("## Data availability", "## Data Availability", 1),
    )
    _add_markdown(
        document,
        figure_legends.replace("## Figure legends", "## Figure Legends", 1),
    )
    figure_names = {
        1: "workflow",
        2: "ecm_paradox",
        3: "msigdb_audit",
        4: "pairwise_concordance",
        5: "sensitivity",
    }
    for index in range(1, 6):
        figure_name = figure_names[index]
        figure_path = FIGURES / f"fig{index}_{figure_name}.png"
        if figure_path.exists():
            document.add_picture(str(figure_path), width=Inches(6.5))
            document.add_paragraph(f"Figure {index}.")
    _add_markdown(document, references)
    document.save(OUTPUT / "JBI_Manuscript.docx")

    submission_figures = OUTPUT / "figures"
    submission_figures.mkdir(parents=True, exist_ok=True)
    for index, figure_name in figure_names.items():
        for suffix in (".png", ".pdf"):
            source_figure = FIGURES / f"fig{index}_{figure_name}{suffix}"
            if source_figure.exists():
                shutil.copy2(
                    source_figure,
                    submission_figures / f"Figure_{index}{suffix}",
                )
    # Figure 6 (positive controls and power calibration) is submitted as
    # supplementary Figure S1 to stay within the 8 display-item limit.
    for suffix in (".png", ".pdf"):
        source_figure = FIGURES / f"fig6_calibration{suffix}"
        if source_figure.exists():
            shutil.copy2(
                source_figure,
                submission_figures / f"Figure_S1{suffix}",
            )

    title_page = f"""# Title Page

**Article type:** Special Communication

**Title:** Module-level agreement is not replication: a reporting and reproducibility gate for public human intervertebral disc transcriptomes

**Authors:** Zhongyuan Liu

**Affiliations:** [Author affiliations and full postal addresses]

**Corresponding author:** Zhongyuan Liu, [affiliation], [email]

**Keywords:** reproducibility; transcriptomics; bioinformatics; gene-sets; intervertebral-disc

**Word count:** {body_word_count:,} words

**Figures:** 5

**Tables:** 3 (Tables 1-2 and Box 1)

**Supplementary material:** Tables S1-S27, Figure S1, methods and sensitivity outputs
"""
    _write_docx(OUTPUT / "JBI_Title_Page.docx", title_page, "Title Page")
    (OUTPUT / "JBI_Title_Page.md").write_text(title_page, encoding="utf-8")

    cover = """# Cover Letter

Dear Editors,

We submit our manuscript, "Module-level agreement is not replication: a reporting and reproducibility gate for public human intervertebral disc transcriptomes," for consideration as a Special Communication.

The manuscript addresses a practical problem in public-data biomedical informatics: module-level significance and direction agreement are often described as replication even when the same genes do not move consistently across cohorts. We performed a matched audit of public human intervertebral disc transcriptome cohorts and operationalized the distinction between exploratory, internally stable, contextually concordant, candidate-replicated and FDR-controlled claims.

The contribution is not a new statistical test or algorithm. It is an integrated methodological audit with a reusable claim ladder and a sensitivity battery spanning donor dependence, probe selection, gene identity, platform scale and FDR-family structure. The results show that module-level significance can coexist with gene-level concordance that does not exceed chance, and that interpretation depends strongly on these design and analysis choices.

The manuscript fits JBI's emphasis on generalizable methodological lessons for developers and users of biomedical informatics methods. The work is original, is not under consideration elsewhere, and all authors have approved the submission.

Sincerely,

Zhongyuan Liu
"""
    _write_docx(OUTPUT / "JBI_Cover_Letter.docx", cover, "Cover Letter")
    (OUTPUT / "JBI_Cover_Letter.md").write_text(cover, encoding="utf-8")

    declarations = """# Declarations

## Competing Interests

The authors declare no competing interests.

## Funding

This research did not receive any specific grant from funding agencies in the public, commercial or not-for-profit sectors.

## CRediT Author Contributions

Zhongyuan Liu: Conceptualization, Methodology, Software, Formal analysis, Validation, Investigation, Data curation, Writing - original draft, Writing - review and editing.

## Declaration of Generative AI and AI-Assisted Technologies

During the preparation of this work, the author(s) used Codex, an OpenAI coding agent, to assist with data-analysis code generation, statistical sensitivity analyses, literature traceability, and manuscript editing. After using this tool, the author(s) reviewed and edited the content as needed and take full responsibility for the content of the published article.

## Data Statement

All source cohorts are public GEO series: GSE70362, GSE23130, GSE186542, GSE167199, GSE146904 and GSE207176. Derived analysis tables and the audit code are available at https://github.com/wdsve/ivd-repro-audit and archived at https://doi.org/10.5281/zenodo.23133983. The manuscript uses no private patient-level data.

## Ethics Statement

This study used publicly available de-identified datasets and did not involve new recruitment or intervention.
"""
    _write_docx(OUTPUT / "JBI_Declarations.docx", declarations, "Declarations")
    (OUTPUT / "JBI_Declarations.md").write_text(declarations, encoding="utf-8")

    supplement_text = SUPPLEMENT.read_text(encoding="utf-8")
    _write_docx(OUTPUT / "JBI_Supplement.docx", supplement_text, "Supplementary Material")

    _graphical_abstract()


if __name__ == "__main__":
    main()
