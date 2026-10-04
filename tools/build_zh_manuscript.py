#!/usr/bin/env python
"""Build a Chinese DOCX version of the manuscript for reading/review.

Reuses the styling and markdown rendering from build_jbi_submission.py and
adds figure embedding plus CJK font handling. The Chinese manuscript is a
translation aid only; the English package in jbi_submission/ is the
submission version.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

import build_jbi_submission as bs

ROOT = TOOLS.parent
SOURCE = ROOT / "manuscript" / "manuscript_zh.md"
OUTPUT = ROOT / "manuscript" / "manuscript_zh.docx"

IMAGE_RE = re.compile(r"^!\[[^\]]*\]\((?P<src>[^)]+)\)\s*$")


def _set_cjk_font(document: Document) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), "宋体")


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.lstrip().splitlines()
    title = "中文稿"
    if lines and lines[0].startswith("# "):
        title = lines[0][2:].strip()
        lines = lines[1:]

    document = Document()
    bs._style_document(document)
    _set_cjk_font(document)
    document.add_heading(title, level=0)

    chunk: list[str] = []

    def flush() -> None:
        if chunk:
            bs._add_markdown(document, "\n".join(chunk))
            chunk.clear()

    for line in lines:
        match = IMAGE_RE.match(line.strip())
        if match:
            flush()
            image = (SOURCE.parent / match.group("src")).resolve()
            if image.exists():
                document.add_picture(str(image), width=Inches(6.3))
            else:
                document.add_paragraph(f"[missing figure: {image}]")
        else:
            chunk.append(line)
    flush()

    document.save(OUTPUT)
    print(f"written: {OUTPUT}")


if __name__ == "__main__":
    main()
